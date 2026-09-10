"""Exercise chart APIs against actual XML and workbooks instead of mocks."""

import io
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def test_scatter_labels_can_be_enabled_removed_and_reopened():
    data = XyChartData()
    series = data.add_series("Points")
    series.add_data_point(3, 4)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, 0, 0, Inches(8), Inches(5), data).chart
    plot = chart.plots[0]
    plot.has_data_labels = False
    assert not plot.has_data_labels
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.number_format = "0.00"
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    plot = Presentation(output).slides[0].shapes[0].chart.plots[0]
    assert plot.has_data_labels
    assert plot.data_labels.show_value
    assert plot.data_labels.number_format == "0.00"
    plot.has_data_labels = False
    assert not plot.has_data_labels


@pytest.mark.parametrize("kind", ["category", "xy"])
def test_chart_workbooks_preserve_formula_like_labels_as_text(kind):
    labels = ("=1+1", "https://example.com", "+3", "@value")
    if kind == "category":
        data = CategoryChartData()
        data.categories = labels
        data.add_series("=SUM(A1:A2)", [1, 2, 3, 4])
    else:
        data = XyChartData()
        for index, label in enumerate(labels):
            data.add_series(label).add_data_point(index, index + 1)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(io.BytesIO(data.xlsx_blob)) as archive:
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        assert sheet.findall(".//s:f", ns) == []
        assert sheet.findall(".//s:hyperlinks", ns) == []
        strings = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        actual = {"".join(item.itertext()) for item in strings}
        assert set(labels).issubset(actual)
