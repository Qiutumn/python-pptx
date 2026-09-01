"""Tests for editable Office Math support."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.math import MathConversionError, latex_to_omml, mathml_to_omml
from pptx.util import Inches, Pt


def test_it_converts_latex_to_inline_and_display_omml():
    inline = latex_to_omml(r"E=mc^2")
    display = latex_to_omml(r"\frac{-b\pm\sqrt{b^2-4ac}}{2a}", display=True)

    assert "<m:oMath" in inline
    assert "<m:sSup>" in inline
    assert "<m:oMathPara" in display
    assert "<m:f>" in display
    assert "<m:rad>" in display


def test_it_converts_mathml_to_omml():
    omml = mathml_to_omml(
        '<math xmlns="http://www.w3.org/1998/Math/MathML">'
        "<msup><mi>x</mi><mn>2</mn></msup></math>"
    )

    assert "<m:oMath" in omml
    assert "<m:sSup>" in omml


def test_it_rejects_unsafe_or_non_math_omml():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text_frame.paragraphs[0]

    with pytest.raises(MathConversionError, match="DOCTYPE"):
        paragraph.add_math('<!DOCTYPE math [<!ENTITY x "x">]><m:oMath><m:t>&x;</m:t></m:oMath>')
    with pytest.raises(MathConversionError, match="root"):
        paragraph.add_math("<a:r><a:t>x</a:t></a:r>")


def test_it_adds_and_round_trips_an_inline_equation():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "The solution is "
    equation = paragraph.add_latex(
        r"x=\frac{-b\pm\sqrt{b^2-4ac}}{2a}",
        font_size=Pt(24),
        color=RGBColor(0x17, 0x3B, 0x57),
    )
    paragraph.add_run().text = "."

    assert len(slide.shapes) == 1
    assert equation.display is False
    assert equation.text.startswith("x=")
    assert len(paragraph.runs) == 2
    assert len(paragraph.math_runs) == 1
    assert paragraph.math_runs[0].element is equation.element

    stream = BytesIO()
    prs.save(stream)
    slide_xml = _slide_xml(stream)
    assert "<mc:AlternateContent" in slide_xml
    assert (
        '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"' in slide_xml
    )
    assert "<a14:m" in slide_xml
    assert "<m:f>" in slide_xml
    assert 'sz="2400"' in slide_xml
    assert '<a:srgbClr val="173B57"' in slide_xml

    stream.seek(0)
    reopened = Presentation(stream)
    reopened_paragraph = reopened.slides[0].shapes[0].text_frame.paragraphs[0]
    assert len(reopened.slides[0].shapes) == 1
    assert reopened_paragraph.text.startswith("The solution is x=")
    assert reopened_paragraph.text.endswith(".")
    assert len(reopened_paragraph.math_runs) == 1
    assert "<m:f>" in reopened_paragraph.math_runs[0].omml


def test_it_adds_a_display_equation_and_removes_it_on_clear():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    paragraph = shape.text_frame.paragraphs[0]

    equation = paragraph.add_latex(r"\sum_{i=1}^{n} i", display=True, font_size=Pt(30))

    assert equation.display is True
    assert "<m:oMathPara" in equation.omml
    assert "<m:nary>" in equation.omml
    paragraph.clear()
    assert paragraph.text == ""
    assert paragraph.math_runs == ()


def test_it_adds_math_inside_a_table_cell():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(5), Inches(1.5))
    paragraph = table_shape.table.cell(0, 0).text_frame.paragraphs[0]

    paragraph.add_latex(r"\alpha+\beta=\gamma")

    assert len(slide.shapes) == 1
    assert slide.shapes[0].has_table
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    reopened = Presentation(stream)
    reopened_cell = reopened.slides[0].shapes[0].table.cell(0, 0)
    assert len(reopened_cell.text_frame.paragraphs[0].math_runs) == 1


def _slide_xml(stream: BytesIO) -> str:
    stream.seek(0)
    with ZipFile(stream) as package:
        return package.read("ppt/slides/slide1.xml").decode("utf-8")
